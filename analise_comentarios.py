"""
Script para análise de comentários de produtos usando OpenAI GPT-4-Turbo
Autor: Análise de Comentários Grupo NC
Data: 2025-10-17
"""

import os
import sys
import json
from pathlib import Path
from dotenv import load_dotenv
from openai import OpenAI
import openpyxl
from openpyxl import load_workbook
import pandas as pd
from time import sleep

# Configurações
ARQUIVO_EXCEL = "Teste Avaliação de Comentários.xlsx"
SHEET_ORIGEM = "Comentários de Produtos"
SHEET_DESTINO = "Resumo de Análise"
MODELO_OPENAI = "gpt-4-turbo"
MAX_CARACTERES = 50


def carregar_configuracao():
    """Carrega variáveis de ambiente e valida configurações"""
    print("📋 Carregando configurações...")
    
    # Carregar .env
    load_dotenv()
    
    # Validar API Key
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key or api_key == "your_api_key_here":
        print("❌ ERRO: OPENAI_API_KEY não encontrada no arquivo .env")
        print("Por favor, adicione sua chave da OpenAI no arquivo .env")
        sys.exit(1)
    
    # Validar arquivo Excel
    if not Path(ARQUIVO_EXCEL).exists():
        print(f"❌ ERRO: Arquivo '{ARQUIVO_EXCEL}' não encontrado")
        sys.exit(1)
    
    print("✅ Configurações carregadas com sucesso")
    return api_key


def ler_comentarios_excel():
    """Lê comentários do Excel e agrupa por produto"""
    print(f"\n📖 Lendo dados de '{SHEET_ORIGEM}'...")
    
    try:
        # Ler Excel
        df = pd.read_excel(ARQUIVO_EXCEL, sheet_name=SHEET_ORIGEM)
        
        # Validar colunas
        if 'Produto' not in df.columns or 'Comentário' not in df.columns:
            print("❌ ERRO: Colunas 'Produto' e 'Comentário' não encontradas")
            sys.exit(1)
        
        # Remover linhas com valores nulos
        df = df.dropna(subset=['Produto', 'Comentário'])
        
        # Agrupar comentários por produto
        comentarios_agrupados = {}
        for produto, grupo in df.groupby('Produto'):
            comentarios = grupo['Comentário'].tolist()
            comentarios_agrupados[produto] = comentarios
        
        total_comentarios = len(df)
        total_produtos = len(comentarios_agrupados)
        
        print(f"✅ Lidos {total_comentarios} comentários de {total_produtos} produtos")
        
        # Mostrar estatísticas
        for produto in sorted(comentarios_agrupados.keys()):
            qtd = len(comentarios_agrupados[produto])
            print(f"   - {produto}: {qtd} comentários")
        
        return comentarios_agrupados
    
    except Exception as e:
        print(f"❌ ERRO ao ler Excel: {e}")
        sys.exit(1)


def analisar_comentarios_com_openai(cliente_openai, produto, comentarios):
    """Analisa comentários de um produto usando OpenAI"""
    
    # Preparar prompt
    comentarios_texto = "\n".join([f"- {com}" for com in comentarios])
    
    prompt = f"""Você receberá comentários sobre um produto. Analise TODOS os comentários e identifique:
1. Os principais pontos POSITIVOS recorrentes (resumo com EXATAMENTE {MAX_CARACTERES} caracteres ou menos)
2. Os principais pontos NEGATIVOS recorrentes (resumo com EXATAMENTE {MAX_CARACTERES} caracteres ou menos)

Se não houver comentários positivos, retorne "Sem pontos positivos".
Se não houver comentários negativos, retorne "Sem pontos negativos".

Comentários do produto:
{comentarios_texto}

Responda APENAS em formato JSON válido, sem nenhum texto adicional:
{{
  "positivos": "texto com máximo {MAX_CARACTERES} caracteres",
  "negativos": "texto com máximo {MAX_CARACTERES} caracteres"
}}"""

    try:
        # Chamar API OpenAI
        response = cliente_openai.chat.completions.create(
            model=MODELO_OPENAI,
            messages=[
                {"role": "system", "content": "Você é um assistente especializado em análise de sentimentos e resumo de avaliações de produtos. Sempre responda em português do Brasil e em formato JSON."},
                {"role": "user", "content": prompt}
            ]
            # temperature removido para compatibilidade com GPT-5
        )
        
        # Extrair resposta
        resposta_texto = response.choices[0].message.content.strip()
        
        # Tentar parsear JSON
        # Remover possíveis markdown code blocks
        if resposta_texto.startswith("```"):
            resposta_texto = resposta_texto.split("```")[1]
            if resposta_texto.startswith("json"):
                resposta_texto = resposta_texto[4:]
            resposta_texto = resposta_texto.strip()
        
        resultado = json.loads(resposta_texto)
        
        # Validar e truncar se necessário
        positivos = resultado.get("positivos", "N/A")[:MAX_CARACTERES]
        negativos = resultado.get("negativos", "N/A")[:MAX_CARACTERES]
        
        return {
            "positivos": positivos,
            "negativos": negativos
        }
    
    except json.JSONDecodeError as e:
        print(f"   ⚠️  Erro ao parsear JSON. Usando valores padrão.")
        return {
            "positivos": "Erro na análise",
            "negativos": "Erro na análise"
        }
    
    except Exception as e:
        print(f"   ⚠️  Erro na API: {e}")
        return {
            "positivos": "Erro na análise",
            "negativos": "Erro na análise"
        }


def analisar_todos_produtos(api_key, comentarios_agrupados):
    """Analisa todos os produtos usando OpenAI"""
    print(f"\n🤖 Iniciando análise com OpenAI ({MODELO_OPENAI})...")
    
    # Inicializar cliente OpenAI
    cliente = OpenAI(api_key=api_key)
    
    resultados = {}
    produtos_ordenados = sorted(comentarios_agrupados.keys())
    total = len(produtos_ordenados)
    
    for idx, produto in enumerate(produtos_ordenados, 1):
        print(f"\n📊 Analisando {produto} ({idx}/{total})...")
        print(f"   Total de comentários: {len(comentarios_agrupados[produto])}")
        
        # Analisar
        resultado = analisar_comentarios_com_openai(
            cliente, 
            produto, 
            comentarios_agrupados[produto]
        )
        
        resultados[produto] = resultado
        
        print(f"   ✅ Positivos: {resultado['positivos']}")
        print(f"   ✅ Negativos: {resultado['negativos']}")
        
        # Pequeno delay para evitar rate limits
        if idx < total:
            sleep(0.5)
    
    print(f"\n✅ Análise concluída para {total} produtos")
    return resultados


def criar_sheet_resultados(resultados):
    """Cria nova sheet com os resultados da análise"""
    print(f"\n📝 Criando sheet '{SHEET_DESTINO}'...")
    
    try:
        # Carregar workbook
        wb = load_workbook(ARQUIVO_EXCEL)
        
        # Remover sheet de destino se já existir
        if SHEET_DESTINO in wb.sheetnames:
            print(f"   ⚠️  Sheet '{SHEET_DESTINO}' já existe. Removendo...")
            del wb[SHEET_DESTINO]
        
        # Criar nova sheet
        ws = wb.create_sheet(SHEET_DESTINO)
        
        # Adicionar cabeçalhos
        ws.append(["Produto", "Resumo dos comentários positivos", "Resumo dos comentários negativos"])
        
        # Formatar cabeçalhos
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        
        # Adicionar dados
        produtos_ordenados = sorted(resultados.keys())
        for produto in produtos_ordenados:
            ws.append([
                produto,
                resultados[produto]["positivos"],
                resultados[produto]["negativos"]
            ])
        
        # Ajustar largura das colunas
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 55
        ws.column_dimensions['C'].width = 55
        
        # Salvar arquivo
        wb.save(ARQUIVO_EXCEL)
        print(f"✅ Sheet '{SHEET_DESTINO}' criada com sucesso!")
        print(f"✅ Arquivo salvo: {ARQUIVO_EXCEL}")
        
    except Exception as e:
        print(f"❌ ERRO ao criar sheet: {e}")
        sys.exit(1)


def main():
    """Função principal"""
    print("=" * 70)
    print("🚀 ANÁLISE DE COMENTÁRIOS DE PRODUTOS - GRUPO NC")
    print("=" * 70)
    
    # 1. Carregar configurações
    api_key = carregar_configuracao()
    
    # 2. Ler comentários do Excel
    comentarios_agrupados = ler_comentarios_excel()
    
    # 3. Analisar com OpenAI
    resultados = analisar_todos_produtos(api_key, comentarios_agrupados)
    
    # 4. Criar sheet com resultados
    criar_sheet_resultados(resultados)
    
    print("\n" + "=" * 70)
    print("✅ PROCESSO CONCLUÍDO COM SUCESSO!")
    print("=" * 70)
    print(f"\nAbra o arquivo '{ARQUIVO_EXCEL}' e veja a sheet '{SHEET_DESTINO}'")


if __name__ == "__main__":
    main()

